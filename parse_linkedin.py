import openpyxl
import csv
import os

files = [
    'LinkedIn-Founders-Owners.xlsx',
    'LinkedIn-Founders-Owners-1.xlsx',
    'LinkedIn-Founders-Owners-2.xlsx',
    'LinkedIn-Founders-Owners_1.xlsx'
]

FOUNDER_KEYWORDS = [
    'founder', 'co-founder', 'cofounder', 'ceo', 'chief executive',
    'owner', 'co-owner', 'president', 'managing director', 'md',
    'cpo', 'chief product', 'cto', 'chief technology', 'entrepreneur'
]

all_leads = []
seen = set()

for f in files:
    if not os.path.exists(f):
        print(f'SKIPPING (not found): {f}')
        continue
    try:
        wb = openpyxl.load_workbook(f)
        ws = wb.active
        headers = [str(cell.value).strip() if cell.value else '' for cell in ws[1]]
        print(f'\nFILE: {f}  |  Rows: {ws.max_row - 1}  |  Headers: {headers}')

        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))
            # Try to find a title/position column
            title_col = None
            for col in headers:
                if any(k in col.lower() for k in ['title', 'position', 'headline', 'job', 'current']):
                    title_col = col
                    break

            title_val = str(row_dict.get(title_col, '') or '').lower()
            is_founder = any(kw in title_val for kw in FOUNDER_KEYWORDS)

            # Build a dedup key from name
            first = str(row_dict.get('First Name', row_dict.get('first name', row_dict.get('FirstName', ''))) or '').strip()
            last = str(row_dict.get('Last Name', row_dict.get('last name', row_dict.get('LastName', ''))) or '').strip()
            email = str(row_dict.get('Email Address', row_dict.get('Email', row_dict.get('email', ''))) or '').strip()
            company = str(row_dict.get('Company', row_dict.get('company', row_dict.get('Organization', ''))) or '').strip()
            title = str(row_dict.get(title_col, '') or '').strip() if title_col else ''
            url = str(row_dict.get('Profile URL', row_dict.get('URL', row_dict.get('LinkedIn URL', row_dict.get('linkedin', ''))) or '')).strip()
            connected_on = str(row_dict.get('Connected On', row_dict.get('connected_on', '')) or '').strip()

            dedup_key = f'{first.lower()}_{last.lower()}_{company.lower()}'
            if dedup_key in seen:
                continue
            seen.add(dedup_key)

            if is_founder or not title_col:  # if no title col found, include all
                all_leads.append({
                    'First Name': first,
                    'Last Name': last,
                    'Title / Headline': title,
                    'Company': company,
                    'Email': email,
                    'LinkedIn URL': url,
                    'Connected On': connected_on,
                    'Source File': f
                })
    except Exception as e:
        print(f'ERROR reading {f}: {e}')

print(f'\n✅ Total leads found: {len(all_leads)}')

# Write to CSV
out_file = 'almostnormal_founder_leads.csv'
if all_leads:
    with open(out_file, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.DictWriter(csvfile, fieldnames=all_leads[0].keys())
        writer.writeheader()
        writer.writerows(all_leads)
    print(f'📁 Saved to: {out_file}')

    # Print preview
    print('\n--- PREVIEW (first 5 leads) ---')
    for lead in all_leads[:5]:
        print(lead)
else:
    print('No leads found. Printing all raw headers for debugging:')
