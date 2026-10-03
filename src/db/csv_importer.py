"""Universal lead importer supporting CSV and Excel (.xlsx) files with smart column and sheet detection."""

import csv
import io
import re
from typing import Dict, List, Optional, Tuple

from src.db.database import Database
from src.db.models import Contact

HEADER_ALIASES = {
    "first_name": ["first_name", "firstname", "first", "fname"],
    "last_name": ["last_name", "lastname", "last", "lname", "surname"],
    "full_name": ["full_name", "fullname", "name", "contact_name", "contactname", "person_name"],
    "company": ["company", "company_name", "companyname", "business", "business_name", "businessname", "organization", "org"],
    "job_title": ["job_title", "jobtitle", "title", "position", "role", "designation", "occupation"],
    "email": ["email", "email_address", "emailaddress", "mail", "contact_email", "contactemail", "email_id", "emailid", "work_email"],
    "linkedin_url": ["linkedin_url", "linkedin", "linkedinurl", "linkedin_profile", "linkedinprofile", "profile_url", "linkedin_link", "person_linkedin_url"],
    "notes": ["notes", "note", "title_match", "comments", "description", "details", "info", "context", "memo"],
    "website": ["website", "url", "web", "website_url", "site", "domain", "company_website", "link"],
}


def clean_key(raw_key: str) -> str:
    """Normalizes header string for fuzzy alias matching."""
    return re.sub(r"[^a-z0-9]", "", str(raw_key or "").lower().strip())


def map_columns(headers: List[str]) -> Dict[str, Optional[str]]:
    """Maps actual CSV/Excel header names to standard Contact field names."""
    field_mapping = {}
    normalized_headers = {clean_key(h): h for h in headers if h}

    for standard_field, aliases in HEADER_ALIASES.items():
        matched_header = None
        for alias in aliases:
            norm_alias = clean_key(alias)
            if norm_alias in normalized_headers:
                matched_header = normalized_headers[norm_alias]
                break
        field_mapping[standard_field] = matched_header

    return field_mapping


def find_header_row_in_matrix(rows: List[List[any]], max_scan: int = 10) -> Tuple[int, List[str]]:
    """Scans the top rows to find the actual header row containing Name, Company, or Email."""
    best_idx = 0
    best_score = 0
    best_headers = []

    for i, row in enumerate(rows[:max_scan]):
        row_str = [str(c or "").strip() for c in row if c is not None]
        cleaned_cells = [clean_key(c) for c in row_str]
        score = 0
        if any(c in cleaned_cells for c in ["name", "fullname", "firstname", "first"]):
            score += 3
        if any(c in cleaned_cells for c in ["company", "companyname", "business"]):
            score += 3
        if any(c in cleaned_cells for c in ["email", "emailaddress", "mail"]):
            score += 3
        if any(c in cleaned_cells for c in ["linkedin", "linkedinurl", "position", "title"]):
            score += 2

        if score > best_score:
            best_score = score
            best_idx = i
            best_headers = row_str

    return best_idx, best_headers


def parse_rows_into_contacts(headers: List[str], data_rows: List[List[any]]) -> Tuple[List[Contact], List[str]]:
    """Converts raw tabular rows into Contact objects with auto-detection."""
    contacts = []
    errors = []
    col_map = map_columns(headers)

    has_any_contact_col = any(col_map.get(k) for k in ["full_name", "first_name", "company", "email", "linkedin_url"])
    if not has_any_contact_col:
        return [], ["No contact fields (Name, Company, Email, or LinkedIn) were found in the table header."]

    fn_col = col_map.get("first_name")
    ln_col = col_map.get("last_name")
    full_col = col_map.get("full_name")
    comp_col = col_map.get("company")
    title_col = col_map.get("job_title")
    email_col = col_map.get("email")
    li_col = col_map.get("linkedin_url")
    notes_col = col_map.get("notes")
    web_col = col_map.get("website")

    header_to_idx = {h: idx for idx, h in enumerate(headers)}

    def get_val(row, col_name):
        if not col_name or col_name not in header_to_idx:
            return ""
        idx = header_to_idx[col_name]
        if idx < len(row) and row[idx] is not None:
            return str(row[idx]).strip()
        return ""

    for row in data_rows:
        # Check if entire row is empty
        if not any(row):
            continue

        raw_email = get_val(row, email_col).lower()
        li_url = get_val(row, li_col)
        comp = get_val(row, comp_col)
        title = get_val(row, title_col)
        notes = get_val(row, notes_col)
        web = get_val(row, web_col)

        # Names
        fn = get_val(row, fn_col)
        ln = get_val(row, ln_col)
        full = get_val(row, full_col)

        if not fn and full:
            parts = full.split(maxsplit=1)
            fn = parts[0]
            ln = parts[1] if len(parts) > 1 else ""

        # Need at least a name or company or email to be a valid prospect
        if not fn and not comp and not raw_email:
            continue

        # If email is missing, generate unique placeholder for LinkedIn tracking
        if not raw_email or "@" not in raw_email:
            ident = li_url or f"{fn}_{comp}"
            raw_email = f"li_{abs(hash(ident)) % 100000000}@linkedin-lead.local"

        contact = Contact(
            first_name=fn or "there",
            last_name=ln,
            company=comp or "Local Business",
            job_title=title,
            linkedin_url=li_url,
            email=raw_email,
            notes=notes,
            website=web,
        )
        contacts.append(contact)

    return contacts, errors


def parse_csv_file(file_bytes: bytes) -> Tuple[List[Contact], List[str], Dict[str, any]]:
    """Parses a CSV file from raw bytes."""
    try:
        content = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            content = file_bytes.decode("latin-1")
        except Exception as e:
            return [], [f"Encoding error reading CSV: {e}"], {}

    reader = list(csv.reader(io.StringIO(content)))
    if not reader:
        return [], ["Uploaded CSV is empty."], {}

    # Check if this is the "Summary" tab of an exported workbook
    first_few_lines = "".join(content.splitlines()[:5]).lower()
    if "filter" in first_few_lines and "breakdown" in content.lower() and "tab a" in content.lower():
        return [], [
            "⚠️ This uploaded file appears to be the 'Summary' sheet of your Excel file, which only contains statistics! "
            "Please upload the full Excel workbook (.xlsx) or export 'Tab A — Niche Fit' as CSV."
        ], {"is_summary_tab": True}

    header_idx, headers = find_header_row_in_matrix(reader)
    data_rows = reader[header_idx + 1 :]
    contacts, errors = parse_rows_into_contacts(headers, data_rows)

    return contacts, errors, {"headers": headers, "header_row": header_idx + 1}


def parse_excel_file(file_bytes: bytes, target_sheet: Optional[str] = None) -> Tuple[List[Contact], List[str], Dict[str, any]]:
    """Parses an Excel (.xlsx) file, returning contacts, errors, and metadata with available sheets."""
    try:
        import openpyxl
    except ImportError:
        return [], ["openpyxl library not installed. Please install it to read .xlsx files."], {}

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    all_sheets = wb.sheetnames

    # If target sheet not specified, auto-detect best candidate (prefer Niche Fit, Contacts, Leads, or skip Summary)
    if not target_sheet or target_sheet not in all_sheets:
        target_sheet = None
        for s in all_sheets:
            s_lower = s.lower()
            if "niche" in s_lower or "fit" in s_lower or "lead" in s_lower or "contact" in s_lower:
                target_sheet = s
                break
        if not target_sheet:
            # Pick first sheet that isn't named "Summary"
            non_summary = [s for s in all_sheets if "summary" not in s.lower()]
            target_sheet = non_summary[0] if non_summary else all_sheets[0]

    ws = wb[target_sheet]
    rows = list(ws.iter_rows(values_only=True))

    header_idx, headers = find_header_row_in_matrix(rows)
    data_rows = rows[header_idx + 1 :]
    contacts, errors = parse_rows_into_contacts(headers, data_rows)

    meta = {
        "all_sheets": all_sheets,
        "selected_sheet": target_sheet,
        "headers": headers,
        "header_row": header_idx + 1,
    }
    return contacts, errors, meta


def parse_any_lead_file(file_bytes: bytes, filename: str, target_sheet: Optional[str] = None) -> Tuple[List[Contact], List[str], Dict[str, any]]:
    """Universal loader for CSV and XLSX files."""
    if filename.lower().endswith(".xlsx") or filename.lower().endswith(".xlsm"):
        return parse_excel_file(file_bytes, target_sheet=target_sheet)
    else:
        return parse_csv_file(file_bytes)


def import_contacts_to_db(db: Database, contacts: List[Contact]) -> int:
    """Inserts a list of parsed Contact objects into the database, returning imported count."""
    count = 0
    for contact in contacts:
        db.insert_contact(contact)
        count += 1
    return count

